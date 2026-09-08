"""
data/ingest_trektellen_crosstabs_xlsx.py
--------------------------------------------
Parses the richest Trektellen source this project has: an "all presets"
year-totals workbook (one sheet per Trektellen dropdown period preset --
each calendar month, "All months", two named seasons, and every rolling
2- and 3-month window), each sheet holding THREE stacked metric blocks
(Combined, Newly ringed, Retraps) x species x year (2017-2026).

Sheet layout (see this module's tests for a worked example): row 1 is a
title, row 2 blank, row 3 is the header (Metric, Species, <years...>,
Total 2017-2026), then repeating blocks of species rows terminated by a
"TOTAL -- <metric>" row (species=None, only the Total column filled).

Output is one long-format DataFrame: species, year, period_code,
period_label, metric, count, source_file -- the same grain the
consolidated database's trektellen_counts table uses (see
scripts/build_consolidated_database.py), so this is the single
richest-resolution Trektellen fact table available today.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd

_METRIC_LABEL_TO_CODE = {
    "Combined (newly ringed + retraps)": "combined",
    "Newly ringed": "newly_ringed",
    "Retraps": "retraps",
}


def read_index(xlsx_path: str | Path) -> pd.DataFrame:
    """
    The workbook's 'Index' sheet: one row per period preset, with
    period_code, period_label, months_included, and summary totals
    (species_recorded, total_combined, total_newly_ringed, total_retraps)
    across all years combined. Useful for sanity-checking the per-sheet
    parse against a known total.
    """
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb["Index"]
    rows = list(ws.iter_rows(values_only=True))
    header_row_idx = next(i for i, r in enumerate(rows) if r and r[0] == "Period code")
    records = []
    for r in rows[header_row_idx + 1 :]:
        if not r or r[0] is None:
            continue
        records.append(
            {
                "period_code": int(r[0]),
                "period_label": r[1],
                "months_included": r[2],
                "species_recorded": int(r[3]),
                "total_combined": int(r[4]),
                "total_newly_ringed": int(r[5]),
                "total_retraps": int(r[6]),
            }
        )
    wb.close()
    return pd.DataFrame(records)


def _parse_sheet(ws, period_code: int, period_label: str, source_file: str) -> list[dict]:
    """
    Each data row carries its metric label directly in column 0 (it is
    NOT a one-time block header -- the same label repeats on every row
    within that block, e.g. every species row under "Combined ..." says
    "Combined ..." in column 0). "TOTAL -- <metric>" rows are block
    subtotals (species=None) and are skipped -- not per-species facts.
    """
    rows = list(ws.iter_rows(values_only=True))
    header_row_idx = next(i for i, r in enumerate(rows) if r and r[0] == "Metric")
    years = [int(y) for y in rows[header_row_idx][2:-1]]  # between "Species" and "Total 2017-2026"

    records = []
    for r in rows[header_row_idx + 1 :]:
        if not r or r[0] is None:
            continue
        label = r[0]
        if not isinstance(label, str) or label not in _METRIC_LABEL_TO_CODE:
            continue  # "TOTAL -- ..." rows and any stray non-data rows
        metric = _METRIC_LABEL_TO_CODE[label]

        species = r[1]
        if species is None:
            continue
        year_values = r[2 : 2 + len(years)]
        for year, count in zip(years, year_values):
            if count is None:
                continue
            records.append(
                {
                    "species": species,
                    "year": year,
                    "period_code": period_code,
                    "period_label": period_label,
                    "metric": metric,
                    "count": float(count),
                    "source_file": source_file,
                }
            )
    return records


def parse_crosstabs_xlsx(xlsx_path: str | Path) -> pd.DataFrame:
    """
    Parses every non-Index sheet into the unified long format described
    in this module's docstring. period_code is looked up from the Index
    sheet by matching period_label to sheet name.
    """
    xlsx_path = Path(xlsx_path)
    index = read_index(xlsx_path)
    label_to_code = dict(zip(index["period_label"], index["period_code"]))

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    all_records = []
    for sheet_name in wb.sheetnames:
        if sheet_name == "Index":
            continue
        period_code = label_to_code.get(sheet_name)
        if period_code is None:
            continue  # sheet name didn't match any Index row -- skip rather than guess a code
        all_records.extend(_parse_sheet(wb[sheet_name], period_code, sheet_name, xlsx_path.name))
    wb.close()

    return pd.DataFrame(all_records)
